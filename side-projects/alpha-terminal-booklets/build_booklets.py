from __future__ import annotations

import math
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Mm
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
ASSETS = Path(__file__).resolve().parent / "assets"
PAGES_ROOT = Path(__file__).resolve().parent / "page-renders"
PDF_ROOT = ROOT / "output" / "pdf"
DOCX_ROOT = ROOT / "output" / "docx"

W, H = 1654, 2339
M = 112
NAVY = "#0B1830"
INK = "#142033"
MUTED = "#637086"
PALE = "#F2F5F8"
LINE = "#D9E0E8"
WHITE = "#FFFFFF"
RED = "#C94747"

FONT_REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_MED = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def hexrgb(value: str):
    value = value.lstrip("#")
    return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))


def wrap(draw: ImageDraw.ImageDraw, text: str, fnt, width: int):
    words = text.split()
    lines, current = [], ""
    for word in words:
        probe = f"{current} {word}".strip()
        if draw.textbbox((0, 0), probe, font=fnt)[2] <= width:
            current = probe
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def text_block(draw, xy, text, fnt, fill, width, leading=None, max_lines=None):
    x, y = xy
    leading = leading or int(fnt.size * 1.32)
    lines = wrap(draw, text, fnt, width)
    if max_lines and len(lines) > max_lines:
        lines = lines[:max_lines]
        tail = lines[-1]
        while draw.textbbox((0, 0), tail + "…", font=fnt)[2] > width and tail:
            tail = tail[:-1]
        lines[-1] = tail.rstrip() + "…"
    for line in lines:
        draw.text((x, y), line, font=fnt, fill=fill)
        y += leading
    return y


def rounded(draw, box, fill, outline=None, radius=22, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def crop_cover(path: Path, size=(W, H)):
    img = Image.open(path).convert("RGB")
    target_ratio = size[0] / size[1]
    ratio = img.width / img.height
    if ratio > target_ratio:
        new_w = int(img.height * target_ratio)
        left = (img.width - new_w) // 2
        img = img.crop((left, 0, left + new_w, img.height))
    else:
        new_h = int(img.width / target_ratio)
        top = (img.height - new_h) // 2
        img = img.crop((0, top, img.width, top + new_h))
    return img.resize(size, Image.Resampling.LANCZOS)


def base_page(book, page_no, section):
    img = Image.new("RGB", (W, H), WHITE)
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, W, 22), fill=book["accent"])
    draw.text((M, 72), book["short"].upper(), font=font(22, True), fill=book["accent"])
    draw.text((W - M, 72), section.upper(), font=font(20, True), fill=MUTED, anchor="ra")
    draw.line((M, 118, W - M, 118), fill=LINE, width=2)
    draw.text((M, H - 72), "ALPHA TERMINAL FIELD GUIDE", font=font(17, True), fill=MUTED)
    draw.text((W - M, H - 72), f"{page_no:02d}", font=font(18, True), fill=book["accent"], anchor="ra")
    return img, draw


def page_title(draw, eyebrow, title, intro, accent, y=160):
    draw.text((M, y), eyebrow.upper(), font=font(22, True), fill=accent)
    y += 48
    y = text_block(draw, (M, y), title, font(57, True), INK, W - 2 * M, leading=68)
    y += 20
    y = text_block(draw, (M, y), intro, font(27), MUTED, W - 2 * M, leading=39)
    return y


def card(draw, box, title, body, accent, tag=None, metric=None):
    x1, y1, x2, y2 = box
    rounded(draw, box, WHITE, LINE, 24, 2)
    draw.rectangle((x1, y1, x1 + 10, y2), fill=accent)
    if tag:
        draw.text((x1 + 34, y1 + 28), tag.upper(), font=font(18, True), fill=accent)
        title_y = y1 + 67
    else:
        title_y = y1 + 30
    draw.text((x1 + 34, title_y), title, font=font(27, True), fill=INK)
    if metric:
        draw.text((x2 - 30, title_y), metric, font=font(25, True), fill=accent, anchor="ra")
    text_block(draw, (x1 + 34, title_y + 48), body, font(21), MUTED, x2 - x1 - 68, leading=31)


def cover_page(book):
    img = crop_cover(ASSETS / book["cover"])
    img = ImageEnhance.Contrast(img).enhance(1.08)
    overlay = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    od = ImageDraw.Draw(overlay)
    for y in range(H):
        alpha = int(205 * max(0, 1 - y / (H * 0.76)))
        od.line((0, y, W, y), fill=(5, 13, 28, alpha))
    od.rectangle((0, H - 410, W, H), fill=(5, 13, 28, 178))
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((M, 105, M + 346, 160), radius=27, fill=book["accent"])
    draw.text((M + 173, 132), "ALPHA TERMINAL FIELD GUIDE", font=font(18, True), fill=WHITE, anchor="mm")
    y = 230
    y = text_block(draw, (M, y), book["title"], font(75, True), WHITE, W - 2 * M, leading=84)
    y += 24
    text_block(draw, (M, y), book["subtitle"], font(31), "#E5EDF7", W - 2 * M, leading=44)
    draw.text((M, H - 292), book["promise"], font=font(25, True), fill=book["accent"])
    text_block(draw, (M, H - 236), book["cover_line"], font(24), WHITE, W - 2 * M, leading=35)
    draw.text((M, H - 85), "RESEARCH SYSTEMS • INDIA • 2026", font=font(18, True), fill="#D8E0EA")
    return img


def purpose_page(book):
    img, draw = base_page(book, 2, "The terminal's job")
    y = page_title(draw, "01 • Purpose", book["purpose_title"], book["purpose_intro"], book["accent"])
    y += 35
    # photographic strip
    photo = crop_cover(ASSETS / book["cover"], (W - 2 * M, 430)).filter(ImageFilter.GaussianBlur(0.2))
    img.paste(photo, (M, y))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((M, y, W - M, y + 430), radius=24, outline=WHITE, width=4)
    y += 480
    draw.text((M, y), "THE DECISION LOOP", font=font(21, True), fill=book["accent"])
    y += 52
    steps = book["loop"]
    gap = 20
    box_w = (W - 2 * M - gap * (len(steps) - 1)) // len(steps)
    for i, (num, title, note) in enumerate(steps):
        x = M + i * (box_w + gap)
        rounded(draw, (x, y, x + box_w, y + 285), PALE, None, 22)
        draw.ellipse((x + 24, y + 24, x + 74, y + 74), fill=book["accent"])
        draw.text((x + 49, y + 49), num, font=font(19, True), fill=WHITE, anchor="mm")
        draw.text((x + 24, y + 100), title, font=font(24, True), fill=INK)
        text_block(draw, (x + 24, y + 144), note, font(19), MUTED, box_w - 48, leading=28)
    y += 335
    rounded(draw, (M, y, W - M, y + 200), NAVY, None, 20)
    draw.text((M + 35, y + 34), "THE PROMISE", font=font(18, True), fill=book["accent"])
    text_block(draw, (M + 35, y + 76), book["promise_detail"], font(25, True), WHITE, W - 2 * M - 70, leading=36)
    return img


def architecture_page(book):
    img, draw = base_page(book, 3, "System architecture")
    y = page_title(draw, "02 • Architecture", "Build an evidence chain, not a signal pile", book["architecture_intro"], book["accent"])
    y += 55
    # pipeline
    flow = book["pipeline"]
    box_w = 238
    gap = 21
    start_x = M
    for i, (label, note) in enumerate(flow):
        x = start_x + i * (box_w + gap)
        rounded(draw, (x, y, x + box_w, y + 210), WHITE, book["accent"] if i == len(flow) - 1 else LINE, 20, 3)
        draw.text((x + 25, y + 26), f"0{i+1}", font=font(19, True), fill=book["accent"])
        draw.text((x + 25, y + 67), label, font=font(23, True), fill=INK)
        text_block(draw, (x + 25, y + 112), note, font(17), MUTED, box_w - 50, leading=25, max_lines=3)
        if i < len(flow) - 1:
            ax = x + box_w + 5
            draw.line((ax, y + 105, ax + gap - 10, y + 105), fill=book["accent"], width=4)
            draw.polygon([(ax + gap - 10, y + 99), (ax + gap - 2, y + 105), (ax + gap - 10, y + 111)], fill=book["accent"])
    y += 275
    draw.text((M, y), "EVIDENCE STACK", font=font(21, True), fill=book["accent"])
    y += 48
    stack_h = 116
    for i, (name, detail, status) in enumerate(book["stack"]):
        yy = y + i * (stack_h + 12)
        rounded(draw, (M, yy, W - M, yy + stack_h), PALE if i % 2 == 0 else WHITE, LINE, 18, 2)
        draw.text((M + 28, yy + 25), name, font=font(24, True), fill=INK)
        text_block(draw, (M + 360, yy + 22), detail, font(20), MUTED, 820, leading=29, max_lines=2)
        badge_color = book["accent"] if status == "Core" else "#78869A"
        rounded(draw, (W - M - 170, yy + 30, W - M - 28, yy + 80), badge_color, None, 24)
        draw.text((W - M - 99, yy + 55), status.upper(), font=font(16, True), fill=WHITE, anchor="mm")
    y += len(book["stack"]) * (stack_h + 12) + 30
    rounded(draw, (M, y, W - M, y + 180), "#FFF8E8", "#F0D79A", 18, 2)
    draw.text((M + 28, y + 28), "DATA CONTRACT", font=font(18, True), fill="#9B6B00")
    text_block(draw, (M + 28, y + 70), book["data_contract"], font(22), INK, W - 2 * M - 56, leading=33)
    return img


def engines_page(book):
    img, draw = base_page(book, 4, "Alpha engines")
    y = page_title(draw, "03 • Research playbooks", "Three engines, three independent reasons to care", book["engines_intro"], book["accent"])
    y += 50
    card_h = 360
    for i, engine in enumerate(book["engines"]):
        yy = y + i * (card_h + 28)
        x1, x2 = M, W - M
        rounded(draw, (x1, yy, x2, yy + card_h), WHITE, LINE, 24, 2)
        draw.rectangle((x1, yy, x1 + 12, yy + card_h), fill=book["accent"])
        draw.text((x1 + 38, yy + 30), engine[0].upper(), font=font(18, True), fill=book["accent"])
        draw.text((x1 + 38, yy + 72), engine[1], font=font(31, True), fill=INK)
        text_block(draw, (x1 + 38, yy + 125), engine[2], font(21), MUTED, 830, leading=31, max_lines=4)
        # mini signal bars
        bx = x2 - 420
        draw.text((bx, yy + 32), "EVIDENCE GATES", font=font(17, True), fill=MUTED)
        for j, (lab, score) in enumerate(engine[3]):
            by = yy + 84 + j * 72
            draw.text((bx, by), lab, font=font(18, True), fill=INK)
            draw.rounded_rectangle((bx, by + 30, bx + 320, by + 45), radius=8, fill="#E5EAF0")
            draw.rounded_rectangle((bx, by + 30, bx + int(320 * score), by + 45), radius=8, fill=book["accent"])
        draw.text((x1 + 38, yy + card_h - 52), f"Invalidate when: {engine[4]}", font=font(18, True), fill=RED)
    return img


def decision_page(book):
    img, draw = base_page(book, 5, "Decision design")
    y = page_title(draw, "04 • Decision card", "Make every idea auditable before it becomes actionable", book["decision_intro"], book["accent"])
    y += 46
    rounded(draw, (M, y, W - M, y + 920), NAVY, None, 28)
    draw.text((M + 42, y + 35), "RESEARCH IDEA • ILLUSTRATIVE", font=font(18, True), fill=book["accent"])
    draw.text((M + 42, y + 82), book["trade_card"][0], font=font(37, True), fill=WHITE)
    text_block(draw, (M + 42, y + 142), book["trade_card"][1], font(22), "#CBD6E5", W - 2 * M - 84, leading=32)
    fields = book["trade_card"][2]
    col_w = (W - 2 * M - 100) // 2
    fy = y + 270
    for i, (label, value) in enumerate(fields):
        col = i % 2
        row = i // 2
        xx = M + 42 + col * (col_w + 30)
        yy = fy + row * 170
        draw.text((xx, yy), label.upper(), font=font(17, True), fill=book["accent"])
        text_block(draw, (xx, yy + 38), value, font(21, True), WHITE, col_w, leading=30, max_lines=3)
    yy = y + 792
    draw.line((M + 42, yy, W - M - 42, yy), fill="#45526A", width=2)
    draw.text((M + 42, yy + 30), "DECISION", font=font(17, True), fill=book["accent"])
    draw.text((M + 200, yy + 26), book["trade_card"][3], font=font(25, True), fill=WHITE)
    y += 970
    draw.text((M, y), "THE FIVE GATES", font=font(21, True), fill=book["accent"])
    y += 52
    gates = ["Fresh", "Independent", "Point-in-time", "Sized", "Falsifiable"]
    gap = 18
    gw = (W - 2 * M - gap * 4) // 5
    for i, gate in enumerate(gates):
        x = M + i * (gw + gap)
        rounded(draw, (x, y, x + gw, y + 150), PALE, LINE, 18, 2)
        draw.ellipse((x + 18, y + 20, x + 50, y + 52), fill=book["accent"])
        draw.text((x + 34, y + 36), "✓", font=font(18, True), fill=WHITE, anchor="mm")
        draw.text((x + 18, y + 78), gate, font=font(20, True), fill=INK)
    y += 205
    rounded(draw, (M, y, W - M, y + 180), "#FFF4F4", "#E7B7B7", 18, 2)
    draw.text((M + 28, y + 26), "NO-GO RULE", font=font(18, True), fill=RED)
    text_block(draw, (M + 28, y + 67), book["no_go"], font(22, True), INK, W - 2 * M - 56, leading=32)
    return img


def validation_page(book):
    img, draw = base_page(book, 6, "Validation")
    y = page_title(draw, "05 • Proof", "Backtest the decision process, not the story", book["validation_intro"], book["accent"])
    y += 50
    # split timeline
    draw.text((M, y), "WALK-FORWARD DESIGN", font=font(21, True), fill=book["accent"])
    y += 58
    segments = [("BUILD", 0.42, "Rules + thresholds"), ("VALIDATE", 0.24, "Tune once"), ("HOLD OUT", 0.34, "No peeking")]
    x = M
    total_w = W - 2 * M
    for idx, (label, pct, note) in enumerate(segments):
        sw = int(total_w * pct) - (8 if idx < 2 else 0)
        fill = book["accent"] if idx == 2 else ("#B7C1CF" if idx == 1 else "#E0E5EB")
        rounded(draw, (x, y, x + sw, y + 140), fill, None, 16)
        draw.text((x + 24, y + 27), label, font=font(22, True), fill=WHITE if idx == 2 else INK)
        draw.text((x + 24, y + 76), note, font=font(18), fill=WHITE if idx == 2 else INK)
        x += sw + 8
    y += 205
    draw.text((M, y), "WHAT THE TEST MUST REPORT", font=font(21, True), fill=book["accent"])
    y += 50
    metrics = book["metrics"]
    gap = 18
    mw = (W - 2 * M - gap * 2) // 3
    for i, (name, why) in enumerate(metrics):
        row, col = divmod(i, 3)
        xx = M + col * (mw + gap)
        yy = y + row * 205
        rounded(draw, (xx, yy, xx + mw, yy + 178), PALE, None, 18)
        draw.text((xx + 24, yy + 24), name, font=font(22, True), fill=INK)
        text_block(draw, (xx + 24, yy + 68), why, font(18), MUTED, mw - 48, leading=27, max_lines=3)
    y += 440
    draw.text((M, y), "BIAS CHECK", font=font(21, True), fill=book["accent"])
    y += 50
    for i, item in enumerate(book["biases"]):
        yy = y + i * 98
        draw.ellipse((M, yy + 4, M + 30, yy + 34), fill=RED)
        draw.text((M + 15, yy + 19), "!", font=font(17, True), fill=WHITE, anchor="mm")
        text_block(draw, (M + 48, yy), item, font(21), INK, W - 2 * M - 48, leading=30, max_lines=2)
    y += len(book["biases"]) * 98 + 15
    rounded(draw, (M, y, W - M, y + 190), NAVY, None, 20)
    draw.text((M + 30, y + 28), "PASS CONDITION", font=font(18, True), fill=book["accent"])
    text_block(draw, (M + 30, y + 70), book["pass_condition"], font(23, True), WHITE, W - 2 * M - 60, leading=34)
    return img


def operations_page(book):
    img, draw = base_page(book, 7, "Operating rhythm")
    y = page_title(draw, "06 • Workflow", "Turn research into a repeatable operating cadence", book["operations_intro"], book["accent"])
    y += 45
    draw.text((M, y), "DAILY → WEEKLY → MONTHLY", font=font(21, True), fill=book["accent"])
    y += 52
    cadence = book["cadence"]
    ch = 250
    for i, (label, action, output) in enumerate(cadence):
        yy = y + i * (ch + 20)
        rounded(draw, (M, yy, W - M, yy + ch), WHITE, LINE, 22, 2)
        rounded(draw, (M + 24, yy + 24, M + 220, yy + 80), book["accent"], None, 26)
        draw.text((M + 122, yy + 52), label.upper(), font=font(18, True), fill=WHITE, anchor="mm")
        draw.text((M + 260, yy + 30), action, font=font(26, True), fill=INK)
        text_block(draw, (M + 260, yy + 78), output, font(20), MUTED, W - 2 * M - 300, leading=29, max_lines=4)
    y += len(cadence) * (ch + 20) + 15
    draw.text((M, y), "90-DAY BUILD ROADMAP", font=font(21, True), fill=book["accent"])
    y += 55
    phases = book["roadmap"]
    pw = (W - 2 * M - 36) // 3
    for i, (phase, title, note) in enumerate(phases):
        x = M + i * (pw + 18)
        rounded(draw, (x, y, x + pw, y + 335), PALE, None, 20)
        draw.text((x + 25, y + 24), phase, font=font(18, True), fill=book["accent"])
        draw.text((x + 25, y + 66), title, font=font(24, True), fill=INK)
        text_block(draw, (x + 25, y + 114), note, font(19), MUTED, pw - 50, leading=28)
    return img


def scorecard_page(book):
    img, draw = base_page(book, 8, "Scorecard")
    y = page_title(draw, "07 • Outcomes", "Judge the terminal by decision quality, not prediction theatre", book["score_intro"], book["accent"])
    y += 50
    draw.text((M, y), "THE SCOREBOARD", font=font(21, True), fill=book["accent"])
    y += 52
    items = book["scorecard"]
    sw = (W - 2 * M - 18) // 2
    for i, (metric, target, note) in enumerate(items):
        row, col = divmod(i, 2)
        xx = M + col * (sw + 18)
        yy = y + row * 220
        rounded(draw, (xx, yy, xx + sw, yy + 195), PALE, None, 20)
        draw.text((xx + 25, yy + 24), metric.upper(), font=font(17, True), fill=MUTED)
        draw.text((xx + 25, yy + 61), target, font=font(30, True), fill=book["accent"])
        text_block(draw, (xx + 25, yy + 110), note, font(18), INK, sw - 50, leading=27, max_lines=2)
    y += math.ceil(len(items) / 2) * 220 + 15
    draw.text((M, y), "NON-NEGOTIABLE GUARDRAILS", font=font(21, True), fill=book["accent"])
    y += 50
    for i, guard in enumerate(book["guardrails"]):
        yy = y + i * 78
        draw.ellipse((M, yy + 2, M + 30, yy + 32), fill=book["accent"])
        draw.text((M + 15, yy + 17), "✓", font=font(16, True), fill=WHITE, anchor="mm")
        draw.text((M + 48, yy), guard, font=font(20, True), fill=INK)
    y += len(book["guardrails"]) * 78 + 25
    rounded(draw, (M, y, W - M, y + 295), NAVY, None, 22)
    draw.text((M + 30, y + 27), "SOURCES AND BOUNDARIES", font=font(18, True), fill=book["accent"])
    text_block(draw, (M + 30, y + 68), book["sources"], font(18), "#DCE5F1", W - 2 * M - 60, leading=27)
    draw.text((M + 30, y + 238), "Research and education only. Alpha is uncertain; losses are possible.", font=font(17, True), fill="#FFCF8A")
    return img


BOOKS = [
    {
        "slug": "01-market-regime-terminal",
        "short": "Market Regime Terminal",
        "title": "BUILDING A\nMARKET REGIME\nTERMINAL",
        "subtitle": "A field guide to turning breadth, volatility, policy and market memory into conditional trade research.",
        "accent": "#00B9D6",
        "cover": "regime-cover.png",
        "promise": "METHOD 01 • CONTEXT BEFORE CANDIDATES",
        "cover_line": "Classify the environment first. Search for alpha only where the market's state makes the hypothesis plausible.",
        "purpose_title": "Know the weather before choosing the vehicle",
        "purpose_intro": "A regime terminal does not predict tomorrow. It describes the state of the market, finds historical analogues and narrows which playbooks deserve attention now.",
        "loop": [("1", "Observe", "Price, breadth, volatility, policy and positioning."), ("2", "Classify", "Label the regime with coverage and confidence."), ("3", "Compare", "Find similar point-in-time historical states."), ("4", "Act", "Apply only playbooks that survived validation.")],
        "promise_detail": "One screen should answer: What state are we in? What changed? What usually followed? Which risks would disprove the view?",
        "architecture_intro": "Every conclusion must retain its observation date, source, revision state and coverage. Missing evidence stays missing.",
        "pipeline": [("Acquire", "Nifty, sectors, breadth, futures, policy"), ("Validate", "Freshness, coverage, conflicts"), ("Classify", "Trend, participation, risk"), ("Compare", "Historical analogues"), ("Decide", "Playbook + risk budget")],
        "stack": [("Price regime", "NIFTY 50/500 trend, drawdown, distance from moving averages.", "Core"), ("Participation", "% above 50/200-day averages and 20-session breadth.", "Core"), ("Risk tone", "Realised volatility, downside concentration and drawdown speed.", "Core"), ("Positioning", "Futures/OI quadrants, liquidity and institutional flows.", "Support"), ("Policy + macro", "RBI decisions, inflation, growth and event vintage.", "Support")],
        "data_contract": "Temporary quantitative floor: 12 June 2000, when NSE NIFTY index futures began. Earlier Indian-market history remains documentary context until consistent, auditable series are available.",
        "engines_intro": "An engine is a testable conditional hypothesis. It earns capital only after coverage, confirmation and downside gates pass.",
        "engines": [
            ("Engine A", "Breadth thrust after stress", "Look for participation expanding after a deep drawdown while volatility stops rising. The edge hypothesis is recovery breadth, not a low index level.", [("Breadth", .83), ("Volatility", .66), ("Trend", .58)], "breadth fades or the index makes a new low on expanding volatility."),
            ("Engine B", "Leadership in a healthy trend", "Rank sectors by relative strength only when the broad index trend and participation are supportive. Prefer persistent leadership over one-day spikes.", [("Trend", .88), ("Relative strength", .78), ("Participation", .72)], "leadership loses benchmark-relative momentum and breadth narrows."),
            ("Engine C", "Defensive rotation", "When participation deteriorates and volatility rises, test low-beta leadership, cash buffers and smaller risk budgets instead of forcing directional conviction.", [("Risk tone", .84), ("Defensives", .69), ("Liquidity", .76)], "breadth repairs, volatility contracts and cyclical leadership broadens."),
        ],
        "decision_intro": "The terminal should produce a research card with conditions, evidence, sizing and invalidation—not a mysterious Buy button.",
        "trade_card": ("Breadth-confirmed recovery basket", "Illustrative research card generated after a broad sell-off; not a live recommendation.", [("STATE", "Recovery candidate • confidence 72% • 5/6 clusters fresh"), ("HYPOTHESIS", "Expanding participation can support 20-session excess return."), ("ENTRY EVIDENCE", "NIFTY 500 >50DMA breadth rises for 5 sessions; vol percentile falls."), ("RISK BUDGET", "0.50% portfolio risk • diversified sector basket • no leverage"), ("BENCHMARK", "NIFTY 50 total return over the same holding window."), ("INVALIDATION", "Breadth below trigger, new closing low or event-data conflict.")], "WATCH • wait for a second independent confirmation"),
        "no_go": "Do not score a regime if core price or breadth history is stale, conflicted or below the declared coverage threshold.",
        "validation_intro": "Regime labels and thresholds are especially vulnerable to hindsight. Recreate only what was knowable on each historical date.",
        "metrics": [("Excess return", "Against the matched benchmark and horizon."), ("Hit rate", "Share of independent windows that were positive."), ("Drawdown", "Worst path, not only final return."), ("Turnover", "How often state changes force action."), ("Coverage", "Rows, clusters and dates truly available."), ("Stability", "Across thresholds, eras and volatility states.")],
        "biases": ["Do not use today's index constituents to claim historical breadth without a survivorship warning.", "Do not use revised macro values before their publication date.", "Do not count overlapping forward windows as independent observations.", "Do not tune the classifier on the same crises used to advertise it."],
        "pass_condition": "Promote a playbook only when held-out excess return survives costs, drawdown remains inside budget and results are not dominated by one crisis.",
        "operations_intro": "A disciplined cadence prevents every market move from becoming a new rule.",
        "cadence": [("Daily", "Refresh the state", "Ingest completed closes; verify freshness; flag changed clusters; do not rewrite thresholds."), ("Weekly", "Review candidates", "Compare sectors and stocks within the active regime; update evidence and invalidation."), ("Monthly", "Audit the model", "Review false positives, coverage gaps, costs and regime transitions; freeze changes before validation.")],
        "roadmap": [("DAYS 1–30", "Trusted data", "Point-in-time price, breadth, volatility and policy registry with source audit."), ("DAYS 31–60", "Classifier", "Transparent scores, transition rules, uncertainty and analogue search."), ("DAYS 61–90", "Playbooks", "Walk-forward tests, research cards, portfolio impact and monitoring.")],
        "score_intro": "A useful terminal reduces unforced errors and makes uncertainty visible.",
        "scorecard": [("Coverage", "> 95% core inputs", "No hidden backfills."), ("Freshness", "Completed sessions", "No intraday/EOD mixing."), ("Stability", "Across market eras", "No crisis-only edge."), ("Drawdown", "Inside declared budget", "Path risk is explicit."), ("Decision latency", "Minutes, not hours", "Evidence is pre-joined."), ("Auditability", "Every label traceable", "Source, date and rule retained.")],
        "guardrails": ["Separate observed facts, derived regime labels and portfolio suggestions.", "Freeze rules before held-out testing and include trading costs.", "Keep documentary pre-2000 stories outside the quantitative classifier.", "Use confidence and coverage—not certainty language."],
        "sources": "Primary candidates: NSE/NSE Indices licensed or permitted histories; RBI policy releases and macro vintages; official exchange derivatives records; audited local EOD histories. Project context: PG-terminal Historical Regimes milestones through v25 and the 12 June 2000 quantitative floor.",
    },
    {
        "slug": "02-catalyst-earnings-terminal",
        "short": "Catalyst and Earnings Terminal",
        "title": "BUILDING A\nCATALYST &\nEARNINGS TERMINAL",
        "subtitle": "A field guide to turning expectations, surprises and price confirmation into event-driven research.",
        "accent": "#F05A71",
        "cover": "catalyst-cover.png",
        "promise": "METHOD 02 • EXPECTATIONS MEET EVIDENCE",
        "cover_line": "Trade the gap between what the market expected and what the company or event actually changed—not the headline alone.",
        "purpose_title": "Measure surprise, then watch how price absorbs it",
        "purpose_intro": "A catalyst terminal timestamps the expectation, the new information and the market response. It separates a good result from a result that was better than priced in.",
        "loop": [("1", "Expect", "Record consensus, trend and pre-event positioning."), ("2", "Observe", "Capture the filing, release and exact timestamp."), ("3", "Confirm", "Measure abnormal price/volume response."), ("4", "Manage", "Use drift, time stops and thesis invalidation.")],
        "promise_detail": "One screen should answer: What changed versus expectations? Was the surprise durable? Did price confirm it? Is the reward still larger than the event risk?",
        "architecture_intro": "Event research fails when timestamps are loose. Store what was known before the event separately from what arrived after it.",
        "pipeline": [("Calendar", "Earnings, policy, corporate events"), ("Expect", "Consensus + pre-trend"), ("Parse", "Actuals, guidance, filing"), ("Confirm", "Gap, volume, relative return"), ("Track", "Drift, decay, invalidation")],
        "stack": [("Event ledger", "Scheduled date, actual release time, source and revision.", "Core"), ("Expectation base", "Comparable results, guidance, estimates or explicit no-consensus state.", "Core"), ("Fundamental delta", "Revenue, margin, EPS, cash flow and guidance change.", "Core"), ("Market response", "Gap, abnormal volume and sector-relative return.", "Core"), ("Context", "Market regime, crowding proxy, liquidity and nearby catalysts.", "Support")],
        "data_contract": "Never infer a surprise when no point-in-time expectation exists. Store the release timestamp, source document and every revision; distinguish reported period from publication date.",
        "engines_intro": "Event alpha is usually conditional on expectations and reaction. The same number can be bullish, neutral or bearish at different prices.",
        "engines": [
            ("Engine A", "Earnings surprise + revision", "Test whether a clean positive surprise, improving guidance and upward forward revisions produce post-event excess return after costs.", [("Surprise", .86), ("Guidance", .74), ("Confirmation", .68)], "guidance reverses, estimates fall or the gap is fully rejected."),
            ("Engine B", "Gap and absorption", "Separate opening emotion from institutional absorption. Look for sustained relative strength, volume and closing quality after the initial gap.", [("Gap quality", .71), ("Volume", .82), ("Close", .77)], "price closes through the event low or abnormal volume disappears."),
            ("Engine C", "Sector read-through", "Use a bellwether result to test peers only when revenue drivers, customer mix and cost exposures genuinely overlap.", [("Driver match", .78), ("Peer response", .65), ("Liquidity", .81)], "the peer's economics or reporting calendar breaks the transmission thesis."),
        ],
        "decision_intro": "A good event card preserves the before/after boundary and refuses to call every beat an opportunity.",
        "trade_card": ("Post-results drift candidate", "Illustrative research card after a verified earnings release; not a live recommendation.", [("EVENT", "Quarterly result • source verified • released after market close"), ("SURPRISE", "Revenue +4% vs expectation; margin +120 bps; guidance raised"), ("PRICE CONFIRMATION", "+3.2% gap; 2.4× volume; closes in top quartile of range"), ("RISK BUDGET", "0.35% portfolio risk • no averaging down • liquidity checked"), ("HORIZON", "5/20-session abnormal return versus sector and NIFTY 50"), ("INVALIDATION", "Event low breaks, revision breadth turns negative or filing changes.")], "RESEARCH • enter only after timestamp and liquidity checks pass"),
        "no_go": "Do not trade a 'surprise' sourced from a revised estimate, an undated screenshot or a filing published after the price move being studied.",
        "validation_intro": "Event studies need exact clocks, clean controls and deconflicted windows.",
        "metrics": [("Abnormal return", "Versus sector and broad benchmark."), ("Drift curve", "1, 5, 20 and 60-session path."), ("Gap retention", "How much of day-one move survives."), ("Drawdown", "Adverse excursion after entry."), ("Liquidity", "Capacity and slippage by turnover bucket."), ("Revision breadth", "Direction and persistence after the event.")],
        "biases": ["Do not backfill consensus estimates after the result is known.", "Do not mix pre-open, intraday and after-close events without separate clocks.", "Do not attribute overlapping macro and company events to one catalyst.", "Do not exclude delisted, suspended or loss-making event outcomes."],
        "pass_condition": "Promote a catalyst only when the timestamp is auditable, abnormal return survives matched controls and the edge persists after realistic gap slippage.",
        "operations_intro": "The calendar drives preparation; evidence drives action.",
        "cadence": [("Daily", "Prepare the event book", "Verify today's releases, timestamps, liquidity and expectation snapshots before the event."), ("Event", "Run the reaction protocol", "Parse filing; compute surprise; observe gap, volume and relative close; issue a card only if gates pass."), ("Weekly", "Review the drift", "Update post-event paths, revisions, thesis decay and time stops; log false positives.")],
        "roadmap": [("DAYS 1–30", "Event ledger", "Official filings, calendars, release clocks and comparable financial schema."), ("DAYS 31–60", "Reaction engine", "Surprise, guidance delta, abnormal return, volume and peer controls."), ("DAYS 61–90", "Drift playbooks", "Walk-forward event studies, trade cards, capacity and monitoring.")],
        "score_intro": "The terminal succeeds when it makes the surprise measurable and the event risk explicit.",
        "scorecard": [("Timestamp quality", "100% verified", "No floating event clocks."), ("Expectation coverage", "Explicit or unavailable", "Never silently inferred."), ("Event attribution", "Deconflicted windows", "Macro overlap tagged."), ("Slippage", "Gap-aware", "No close-to-close fantasy."), ("Drift stability", "Across sectors + years", "Not one earnings season."), ("Thesis decay", "Time stop tracked", "Capital is released.")],
        "guardrails": ["Use official filings and preserve original publication times.", "Separate reported facts, expectation deltas and market interpretation.", "Do not treat a positive result as a positive surprise by default.", "Size for gap risk and avoid leverage around binary events."],
        "sources": "Primary candidates: NSE/BSE corporate filings and official company releases; reviewed macro calendars; point-in-time estimate feeds where licensed; audited price and volume histories. PG-terminal already contains an earnings import foundation and an official-event calendar framework.",
    },
    {
        "slug": "03-relative-strength-portfolio-terminal",
        "short": "Relative Strength Portfolio Terminal",
        "title": "BUILDING A\nRELATIVE STRENGTH\nPORTFOLIO TERMINAL",
        "subtitle": "A field guide to finding leadership, controlling overlap and allocating risk where evidence is strongest.",
        "accent": "#19A875",
        "cover": "relative-strength-cover.png",
        "promise": "METHOD 03 • LEADERSHIP WITH PORTFOLIO CONTEXT",
        "cover_line": "Rank what is leading, test whether participation confirms it, then ask what the candidate changes inside the whole portfolio.",
        "purpose_title": "Turn a watchlist into a portfolio-aware opportunity map",
        "purpose_intro": "A relative-strength terminal compares sectors, stocks and funds on common horizons, then filters leadership through diversification, liquidity and regime fit.",
        "loop": [("1", "Rank", "Compare sectors and securities to the same benchmark."), ("2", "Confirm", "Demand breadth, persistence and liquidity."), ("3", "Simulate", "Measure overlap, beta and concentration impact."), ("4", "Allocate", "Use risk budgets, turnover limits and review dates.")],
        "promise_detail": "One screen should answer: Where is leadership broadening? What exposure do I already own? What does this idea add? What risk must be removed to fund it?",
        "architecture_intro": "A candidate is not attractive in isolation. The engine must see direct holdings, fund look-through and the portfolio's existing factor concentrations.",
        "pipeline": [("Universe", "Sectors, stocks, funds"), ("Rank", "Relative return + persistence"), ("Confirm", "Breadth, trend, liquidity"), ("Simulate", "Overlap, beta, drawdown"), ("Allocate", "Risk + turnover budget")],
        "stack": [("Leadership", "20/60/120-session excess return and ranking persistence.", "Core"), ("Participation", "Sector breadth, new highs and moving-average confirmation.", "Core"), ("Portfolio map", "Direct holdings plus mutual-fund look-through by ISIN.", "Core"), ("Risk map", "Beta, correlation, drawdown, liquidity and concentration.", "Core"), ("Fundamentals", "Valuation, earnings direction and event confirmation.", "Support")],
        "data_contract": "Compare every candidate on the same as-of date, return window and benchmark. Label holdings coverage; unsupported funds and unclassified securities must not masquerade as diversification.",
        "engines_intro": "Leadership can persist, but crowded or duplicated leadership can still reduce portfolio quality.",
        "engines": [
            ("Engine A", "Breadth-confirmed sector leadership", "Rank sector excess returns, then require a growing share of constituents above trend. Prefer broad leadership to index moves driven by a few names.", [("Relative strength", .88), ("Breadth", .81), ("Persistence", .75)], "sector breadth rolls over or excess return breaks its trend."),
            ("Engine B", "Under-owned improving exposure", "Find improving sectors with low current portfolio weight and acceptable correlation. The edge is marginal diversification plus improving evidence.", [("Improvement", .77), ("Low overlap", .84), ("Regime fit", .68)], "correlation spikes, regime fit falls or improvement is not confirmed."),
            ("Engine C", "Redundancy removal", "Identify funds and direct stocks that repeat the same underlying exposures. Reallocate risk budget only after costs, taxes and role clarity are reviewed.", [("Coverage", .79), ("Overlap", .86), ("Cost benefit", .63)], "look-through coverage is incomplete or the duplicate serves a distinct mandate."),
        ],
        "decision_intro": "A portfolio-aware research card shows both expected opportunity and marginal portfolio damage.",
        "trade_card": ("Sector leadership rebalance", "Illustrative research card for a portfolio with duplicated large-cap exposure; not a live recommendation.", [("CANDIDATE", "Top-decile 60-session sector RS • breadth improving 3 weeks"), ("CURRENT EXPOSURE", "6% direct + 9% look-through; 42% overlaps existing leaders"), ("PORTFOLIO IMPACT", "+2% target weight • beta +0.03 • concentration +0.4 pts"), ("RISK BUDGET", "Fund by trimming a redundant exposure; turnover cap applies"), ("BENCHMARK", "Sector excess return vs NIFTY 50 and matched risk bucket"), ("INVALIDATION", "RS rank below median, breadth reversal or portfolio beta breach.")], "SIMULATE • proceed only if marginal diversification stays positive"),
        "no_go": "Do not add a leader when the portfolio-impact simulation is unavailable, holdings coverage is weak or the trade breaches concentration and liquidity limits.",
        "validation_intro": "Rank-based strategies can look excellent before costs and universe bias. Test the complete investable list available on each date.",
        "metrics": [("Excess return", "Against benchmark and risk bucket."), ("Rank persistence", "How long leadership survives."), ("Breadth", "Share of names confirming the move."), ("Turnover", "Rebalance frequency and cost."), ("Marginal risk", "Beta, correlation and concentration change."), ("Capacity", "Liquidity-adjusted deployable size.")],
        "biases": ["Do not rank today's survivors across historical periods.", "Do not compare funds or stocks with mismatched dates and horizons.", "Do not treat unknown fund holdings as zero overlap.", "Do not ignore taxes, exit loads and market impact in rebalance tests."],
        "pass_condition": "Promote a ranking rule only when held-out excess return survives turnover, breadth adds information and portfolio-level drawdown does not worsen beyond budget.",
        "operations_intro": "Leadership moves continuously; portfolio roles should change deliberately.",
        "cadence": [("Daily", "Refresh leadership", "Update completed-close ranks, breadth and alerts; do not rebalance on every rank change."), ("Weekly", "Run portfolio what-if", "Review top candidates, overlap, sector/cap exposure and marginal risk before funding ideas."), ("Monthly", "Rebalance by rule", "Apply rank, risk and turnover thresholds; record what was added, removed and why.")],
        "roadmap": [("DAYS 1–30", "Portfolio map", "Unified holdings upload, security identity, fund look-through and coverage."), ("DAYS 31–60", "Ranking engine", "Common-window returns, breadth, persistence, valuation and regime filters."), ("DAYS 61–90", "Allocator", "What-if cards, risk budgets, turnover, attribution and monitoring.")],
        "score_intro": "The terminal should improve portfolio efficiency, not merely generate more names.",
        "scorecard": [("Holdings coverage", "> 90% value", "Unknown exposure stays visible."), ("Effective positions", "Diversified, not counted", "Overlap-adjusted breadth."), ("Marginal beta", "Inside risk budget", "Every addition simulated."), ("Turnover", "Below declared cap", "Costs remain realistic."), ("Rank stability", "Multiple horizons", "No one-day leadership."), ("Attribution", "Position + sector", "Know what created excess return.")],
        "guardrails": ["Use exact security identity and dated fund holdings for look-through.", "Rank on common dates, horizons and benchmarks.", "Treat overlap as evidence, not an automatic sell instruction.", "Show how every candidate changes concentration, beta and drawdown risk."],
        "sources": "Primary candidates: permitted NSE/NSE Indices histories and classifications; official AMC portfolio disclosures; AMFI identifiers and NAV histories; audited holdings uploads. PG-terminal already has universal holdings analysis, ISIN-normalized overlap, market-cap classification and a regime-fit diversification score.",
    },
]


def save_pdf(pages: list[Path], path: Path):
    c = canvas.Canvas(str(path), pagesize=A4)
    pw, ph = A4
    for page in pages:
        c.drawImage(str(page), 0, 0, width=pw, height=ph)
        c.showPage()
    c.save()


def save_docx(pages: list[Path], path: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(0)
    sec.bottom_margin = Mm(0)
    sec.left_margin = Mm(0)
    sec.right_margin = Mm(0)
    for idx, page in enumerate(pages):
        if idx:
            doc.add_section(WD_SECTION.NEW_PAGE)
            sec = doc.sections[-1]
            sec.page_width = Mm(210)
            sec.page_height = Mm(297)
            sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(0)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = p.paragraph_format.space_after = 0
        p.add_run().add_picture(str(page), width=Mm(210), height=Mm(297))
    doc.save(path)


def main():
    PDF_ROOT.mkdir(parents=True, exist_ok=True)
    DOCX_ROOT.mkdir(parents=True, exist_ok=True)
    PAGES_ROOT.mkdir(parents=True, exist_ok=True)
    for book in BOOKS:
        out_dir = PAGES_ROOT / book["slug"]
        out_dir.mkdir(parents=True, exist_ok=True)
        rendered = [
            cover_page(book), purpose_page(book), architecture_page(book), engines_page(book),
            decision_page(book), validation_page(book), operations_page(book), scorecard_page(book),
        ]
        paths = []
        for i, image in enumerate(rendered, 1):
            path = out_dir / f"page-{i:02d}.png"
            image.save(path, quality=95)
            paths.append(path)
        save_pdf(paths, PDF_ROOT / f"{book['slug']}.pdf")
        save_docx(paths, DOCX_ROOT / f"{book['slug']}.docx")
        print(f"built {book['slug']} ({len(paths)} pages)")


if __name__ == "__main__":
    main()
